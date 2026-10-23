import io 
import magic
from typing import Tuple, Optional

import pdfplumber
from docx import Document 
import PyPDF2

from backend.utils.file_utils import (
    FileParsingError,
    TextExtractionError, 
    FileUploadError,
    log_error,
    log_warning,
    log_info,
    with_fallback
)


from backend.core.config import (
    MAX_FILE_SIZE_MB,
    MAX_FILE_SIZE_BYTES,
    SUPPPORTED_MIME_TYPES
)

class FileParsingError(Exception):
    pass

class FileValidationError(Exception):
    pass

#Validating the uploaded file - file size & file type 
def validate_file(file_data: bytes, file_name: str) -> Tuple[bool, str, Optional[str]]:
    file_size_bytes = len(file_data)
    
    #validating file size
    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = file_size_bytes / (1024*1024)
        return False , (
            f"File size ({size_mb:.2f} MB) exceeds the maximum of {MAX_FILE_SIZE_MB}"
            "Please upload a smalled file or compress your resume."
        ), None
    
    #Validating empty file 
    if file_size_bytes==0: 
        return False, "Upload file is empty ...please check the file you have uploaded and try again.",None

    #Checks the file type internally through its contents not just file name or extension 
    try:
        mime_type = magic.from_buffer(file_data, mime=True)
    except Exception as exc: 
        return False, f"Error determining the file type: {exc}",None

    #Checks whether mime_type is listed in supported file types
    if mime_type not in SUPPPORTED_MIME_TYPES:
        supported=".".join(SUPPPORTED_MIME_TYPES.keys()).upper()

        return False, (
            f"Unsupported file type: {mime_type}"
            f"Please upload one of: {supported}"
        ), None

    return True, "", SUPPPORTED_MIME_TYPES[mime_type]

#######################
#   PDF EXTRACTION   #
#######################

#Helper for extracting hyperlinks 
def extract_pdf_hyperlinks(file_data: bytes) -> str:
    urls = []
    try: 
        reader = PyPDF2.PdfReader(io.BytesIO(file_data))
        for page in reader.pages:
            if "/Annots" not in page:
                continue
            #collects URLs from link annotations
            for annot_ref in page["/Annots"]:
                try: 
                    annot = annot_ref.get_objects()
                    if annot.get("/Subtype") != "/Link":
                        continue
                    action = annot.get("/A", {})
                    uri = action.get("/URI", "")
                    if uri and isinstance(uri, (str, bytes)):
                        #PYPDF2 may return bytes for URI values
                        if isinstance(uri, bytes):
                            uri = uri.decode("urf-8", errors="ignore")
                        uri = uri.strip()
                        if uri.startswith("http"):
                            urls.append(uri)
                except Exception:
                    pass
    except Exception:
        pass
    return "/n".join(urls)

#Extracting PDF with PDFPlumber (main function)
def extract_pdf_with_pdfplumber(file_data: bytes) -> str: 
    text =""
    #Extracting text from each page 
    with pdfplumber.open(io.BytesIO(file_data)) as pdf: 
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "/n"

    #If PDF empty raise error 
    if not text.strip():
        raise TextExtractionError(
            "pdfplumber extracted no text",
            user_message="No text could be extracted from the PDF."
        )

    #get any hyperlinks
    hyperlinks= extract_pdf_hyperlinks(file_data)

    #if any hyperlinks in resume then add it to the text extracted
    if hyperlinks:
        text = text.strip() + "/n" + hyperlinks

    #return extracted text 
    return text.strip()

#Extarcting pdf with PyPDF2 (fallback function)
def extract_pdf_with_pypdf2(file_data: bytes) -> str:
    text = ""
    #Extracting data from each page 
    pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_data))
    for page in pdf_reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text  + "n"

    #If PDF empty raise error 
    if not text.strip():
        raise TextExtractionError(
            "PyPDF2 extract no text",
            user_message="No text could be extracted from the PDF."
        )
    
    #get any hyperlinks
    hyperlinks = extract_pdf_hyperlinks(file_data)
    #if any hyperlinks in resume then add it to the text extracted
    if hyperlinks:
            text = text.strip() + "/n" + hyperlinks

    #return extracted text 
    return text.strip()

#Choose a PDF reader      
def extract_text_from_pdf(file_data:bytes) -> str:
    try:   
        result, used_fallback = with_fallback (
            extract_pdf_with_pdfplumber,
            extract_pdf_with_pypdf2,
            file_data,
            log_fallback=True
        )

        if used_fallback:
            log_info("PDF EXTRACTION succeded using the PyPDF2 fallback", context='resume_parser')
            return result
    #if both functions not working 
    except Exception as e:
        log_error(e, context="extract_text_from_pdf" )
        raise FileParsingError(
            "Failed to extract text from PDF using both pdfplumber and PYPDF2, "
            "The PDF may be corruped, password-protected, or contain only scanned images"
            "Please ensure it contains selectable text."
        )  from e


#######################
#   DOCX EXTRACTION   #
#######################

def extract_text_from_docx(file_data: bytes) -> str:
    try: 
        doc = Document(io.BytesIO(file_data))
        text_parts= []

        #collects text from paragraph
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        #collects text from table cells 
        #Doc stores resume sections in tables therefore need to extract that
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = "/n".join(text_parts)

        if not text.strip():
            raise FileParsingError(
                "No text could be extracted from the document."
                "The document may be empty or corrupted."
            )

        #checks for any hyperlink, if so appends to text
        try: 
            for rel in doc.part.rels.values():
                if "hyperlink" in rel.reltypes.lower():
                    url = rel._target
                    if isinstance(url, str) and url.startswith("http"):
                        text += "/n" + url
        except Exception:
            pass

        log_info(f"Extracted {len(text)} chars from DOCX", context="resume_parser")

        #return extracted text
        return text.strip()

    except FileParsingError:
        raise #Re-raise unchanged 

    except Exception as e:
        log_error(e, context="extract_text_from_docx")
        raise FileParsingError(
            "Failed to extract text from DOCX"
            "The document may be corrupted or in an unsupported format."
            "Please try re-saving or converting to PDF."
        ) from e

##################################################
#   DOC EXTRACTION - NOT SUPPORTED, RAISE ERROR  # 
##################################################

def extract_text_from_doc(file_data: bytes) -> str:
    raise FileParsingError(
        "Legacy .doc format is not supported."
        "Please convert your document to .docx or .pdf and try again."
        "You can convert using Microsoft Word, Google Docs, Or online tools."
    )

#########################################
#   ORCHESTRATOR for file type     #
#########################################

#Chooses an extractor by file type
def extract_text(file_data: bytes, file_type: str) -> str:

    #dispatches to the matching function 
    if file_type=="pdf":
        return extract_text_from_pdf(file_data)
    elif file_type == "docx":
        return extract_text_from_docx(file_data)
    elif file_type== "doc":
        return extract_text_from_doc(file_data)
    else:
        raise FileValidationError(
            f"Inavlid file type: {file_type}. supported types are: pdf, docx and doc"
        )
    
#####################
#   Main Function   #
#####################

def parse_resume_file(file_data: bytes, file_name: str) -> Tuple[str, dict]:
    #Log that parsing has started
    log_info(f"Parsing file: {file_name}", context="parse_resume_file")

    #Stage 1: Validate File
    try:
        is_valid, error_msg, file_type= validate_file(file_data, file_name)
        if  not is_valid:
            log_warning(f"Validation failed for file {file_name}", context="parse_resume_file")
            raise FileValidationError(error_msg)

    except FileValidationError as e:
        raise  #Re-raise the original message 

    #creates new validaion error for unexpected error 
    except Exception as e: 
        log_error(e, context="prase_resume_file_validation")
        raise FileValidationError(
            "Could not validate the uploaded file. Please ensure it is a valid PDF or DOCX."
        ) from e

    #Stage 2: Extraction of file 
    try:
        text = extract_text(file_data, file_type)
        #successful extraction is logged.
        log_info(f"Extracted {len(text)} chars from {file_name}", context="parse_resume_file")

    except FileParsingError:
        raise #Re-raise the original message 
    
    #creates new parsing error for unexpected error 
    except Exception as e:
        log_error(e, context="parse_resume_file_extraction")
        raise FileParsingError(
            "An unexpected error occurred while processing the file."
            "Please try again or contact support if the problem persists."
        ) from e

    metadata = {
        "filename": file_name,
        "file_type": file_type,
        "file_size_bytes": len(file_data),
        "text_length": len(text),
        "success": True,
    }

    return text, metadata


